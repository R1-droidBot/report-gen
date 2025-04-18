import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from PIL import Image, ImageTk
import json
import os
import shutil

EVENTS_FILE = "events.json"
IMAGES_FOLDER = "event_images"
LOGO_PATH = r"C:\Users\Harsh\Downloads\ChatGPT Image Apr 18, 2025, 09_50_26 PM.png"
ICON_PATH = ""

os.makedirs(IMAGES_FOLDER, exist_ok=True)


def load_events():
    if os.path.exists(EVENTS_FILE):
        with open(EVENTS_FILE, "r") as file:
            return json.load(file)
    return []


def save_events(events):
    with open(EVENTS_FILE, "w") as file:
        json.dump(events, file, indent=4)


def add_event():
    event_data = {label: var.get() for label, var in zip(labels, entry_vars)}

    event_data["Objective"] = objective_text.get("1.0", tk.END).strip()
    event_data["Outcome"] = outcome_text.get("1.0", tk.END).strip()
    event_data["Course Contents"] = course_contents_text.get("1.0", tk.END).strip()

    event_data["images"] = []
    for img_path in image_paths:
        if img_path:
            dest_path = os.path.join(IMAGES_FOLDER, os.path.basename(img_path))
            shutil.copy(img_path, dest_path)
            event_data["images"].append(dest_path)

    def copy_special(var_name, var_path):
        if var_path.get():
            dest = os.path.join(IMAGES_FOLDER, os.path.basename(var_path.get()))
            shutil.copy(var_path.get(), dest)
            event_data[var_name] = dest
        else:
            event_data[var_name] = ""

    copy_special("Invitation Letter", invitation_path)
    copy_special("Permission Document", permission_path)
    copy_special("Certificate", certificate_path)
    copy_special("CO-PO Mapping", co_po_path)
    copy_special("Leaflet", leaflet_path)

    if not event_data["Event Number"] or not event_data["Event Name"]:
        messagebox.showerror("Error", "Event Number and Event Name are required!")
        return

    events = load_events()
    if any(e["Event Number"] == event_data["Event Number"] for e in events):
        messagebox.showerror("Error", "Event with this number already exists!")
        return

    events.append(event_data)
    save_events(events)
    messagebox.showinfo("Success", "Event added successfully!")
    clear_fields()
    update_event_list()


def clear_fields():
    for var in entry_vars:
        var.set("")
    objective_text.delete("1.0", tk.END)
    outcome_text.delete("1.0", tk.END)
    course_contents_text.delete("1.0", tk.END)
    image_paths.clear()
    update_image_preview()
    invitation_path.set("")
    permission_path.set("")
    certificate_path.set("")
    co_po_path.set("")
    leaflet_path.set("")
    update_file_labels()


def update_event_list(search_term=None):
    for row in tree.get_children():
        tree.delete(row)

    events = load_events()
    for event in events:
        if search_term:
            search_term = search_term.lower()
            if (search_term in event["Event Number"].lower() or
                    search_term in event["Event Name"].lower() or
                    search_term in event["Date"].lower() or
                    search_term in event["Resource Person"].lower() or
                    search_term in event["Event Type"].lower()):
                tree.insert("", "end", values=(
                    event["Event Number"],
                    event["Event Name"],
                    event["Date"],
                    event["Resource Person"],
                    event["Event Type"],
                    f"{len(event.get('images', []))} images"
                ))
        else:
            tree.insert("", "end", values=(
                event["Event Number"],
                event["Event Name"],
                event["Date"],
                event["Resource Person"],
                event["Event Type"],
                f"{len(event.get('images', []))} images"
            ))


def generate_report():
    selected_item = tree.selection()
    if not selected_item:
        messagebox.showerror("Error", "Please select an event to generate a report.")
        return

    selected_event_number = tree.item(selected_item, "values")[0]
    events = load_events()
    selected_event = next((e for e in events if e["Event Number"] == selected_event_number), None)

    if not selected_event:
        messagebox.showerror("Error", "Event not found.")
        return

    template_path = "event_template.docx"
    if not os.path.exists(template_path):
        messagebox.showerror("Error", "Template file not found!")
        return

    try:
        doc = DocxTemplate(template_path)
        context = {
            'event': {key.replace(" ", "_").lower(): selected_event.get(key, "") for key in selected_event if key not in ["images"]}
        }
        context['document_images'] = {}

        def create_inline_image(path, doc_obj, width_mm=None, height_mm=None):
            if path and os.path.exists(path) and os.path.splitext(path)[1].lower() in ('.jpg', '.jpeg', '.png', '.gif'):
                try:
                    if width_mm is not None and height_mm is not None:
                        return InlineImage(doc_obj, path, width=Mm(width_mm), height=Mm(height_mm))
                    elif width_mm is not None:
                        return InlineImage(doc_obj, path, width=Mm(width_mm))
                    elif height_mm is not None:
                        return InlineImage(doc_obj, path, height=Mm(height_mm))
                    else:
                        return InlineImage(doc_obj, path)
                except Exception as e:
                    print(f"Error creating inline image for {path}: {e}")
            return path  # Return the path if it's not a recognized image

        context['document_images']['invitation'] = create_inline_image(selected_event.get('Invitation Letter', ''), doc, 80)
        context['document_images']['permission'] = create_inline_image(selected_event.get('Permission Document', ''), doc, 80)
        context['document_images']['certificate'] = create_inline_image(selected_event.get('Certificate', ''), doc, 80)
        context['document_images']['co_po'] = create_inline_image(selected_event.get('CO-PO Mapping', ''), doc, 80)
        context['document_images']['leaflet'] = create_inline_image(selected_event.get('Leaflet', ''), doc, 80)

        if "images" in selected_event and selected_event["images"]:
            context['images'] = []
            for idx, img_path in enumerate(selected_event["images"]):
                if os.path.exists(img_path):
                    try:
                        context['images'].append(
                            InlineImage(doc, img_path, width=Mm(50))
                        )
                        context[f'image_{idx}'] = context['images'][-1]
                    except Exception as img_error:
                        print(f"Skipping image {img_path}: {img_error}")
                        continue

        report_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Documents", "*.docx")]
        )
        if report_path:
            doc.render(context)
            doc.save(report_path)
            messagebox.showinfo("Success", "Report generated successfully!")

    except Exception as e:
        messagebox.showerror("Error", f"Report generation failed: {str(e)}")


def delete_event():
    selected_item = tree.selection()
    if not selected_item:
        messagebox.showerror("Error", "Please select an event to delete.")
        return

    selected_event_number = tree.item(selected_item, "values")[0]

    if messagebox.askyesno("Confirm Delete", f"Are you sure you want to delete event {selected_event_number}?"):
        events = load_events()
        events = [e for e in events if e["Event Number"] != selected_event_number]
        save_events(events)
        update_event_list()
        messagebox.showinfo("Success", "Event deleted successfully!")


def search_events():
    search_term = search_var.get()
    update_event_list(search_term)


def add_images():
    files = filedialog.askopenfilenames(
        title="Select Images",
        filetypes=[("Image Files", "*.jpg *.jpeg *.png *.gif")]
    )
    if files:
        image_paths.extend(files)
        update_image_preview()


def update_image_preview():
    for widget in image_preview_frame.winfo_children():
        widget.destroy()

    if not image_paths:
        empty_label = ttk.Label(image_preview_frame, text="No images selected", foreground="gray")
        empty_label.pack(pady=20)
        return

    for i, img_path in enumerate(image_paths):
        try:
            img = Image.open(img_path)
            img.thumbnail((100, 100))
            photo = ImageTk.PhotoImage(img)
            frame = ttk.Frame(image_preview_frame)
            frame.pack(side='left', padx=5, pady=5)

            label = ttk.Label(frame, image=photo)
            label.image = photo
            label.pack()

            # Add a small label with the filename
            filename = os.path.basename(img_path)
            ttk.Label(frame, text=filename, width=15, wraplength=100).pack()
        except Exception as e:
            print(f"Error loading image preview: {e}")


def upload_special_image(var, title):
    file = filedialog.askopenfilename(
        title=f"Select {title}",
        filetypes=[("Image Files", "*.jpg *.jpeg *.png *.gif *.pdf *.docx")]
    )
    if file:
        var.set(file)
        update_file_labels()


def update_file_labels():
    for label, var in file_labels.items():
        filename = os.path.basename(var.get()) if var.get() else "No file selected"
        label.config(text=f"Selected: {filename}" if var.get() else "No file selected")


# Main window setup
root = tk.Tk()
root.title("Event Report Generator")
root.geometry("1200x800")
root.columnconfigure(0, weight=1)
root.rowconfigure(0, weight=1)

# Set the window icon (optional)
if os.path.exists(ICON_PATH):
    try:
        root.iconbitmap(ICON_PATH)
    except tk.TclError as e:
        print(f"Error setting icon: {e}")

image_paths = []

# Create notebook for tabs
notebook = ttk.Notebook(root)
notebook.pack(fill='both', expand=True)

# Create tabs
input_tab = ttk.Frame(notebook)
list_tab = ttk.Frame(notebook)
upload_tab = ttk.Frame(notebook)

notebook.add(input_tab, text="Add Event")
notebook.add(list_tab, text="Event List")
notebook.add(upload_tab, text="Documents & Images")

# Input Tab
main_frame = ttk.Frame(input_tab)
main_frame.pack(fill='both', expand=True, padx=15, pady=15)

# Load the logo image and display it here, after main_frame is defined
try:
    logo_img = Image.open(LOGO_PATH)
    logo_img = logo_img.resize((100, 100))
    logo_photo = ImageTk.PhotoImage(logo_img)
    logo_label = ttk.Label(main_frame, image=logo_photo)
    logo_label.grid(row=0, column=3, rowspan=2, padx=10, pady=10, sticky='nsew')
except FileNotFoundError:
    logo_photo = None
    print(f"Warning: Logo file not found at {LOGO_PATH}")
except Exception as e:
    logo_photo = None
    print(f"Error loading logo: {e}")

# Use a grid layout
main_frame.columnconfigure(0, weight=1)
main_frame.columnconfigure(1, weight=1)
main_frame.columnconfigure(2, weight=1)
if logo_photo:
    main_frame.columnconfigure(3, weight=0)

labels = ["Event Number", "Attendees", "Program Director", "Event Name", "Resource Person",
          "Event IC", "Designation", "Attendance Check", "Date", "Address",
          "Permission Docs", "Event Type", "Funding", "CO-PO Link", "Report Doc",
          "Days", "Remarks", "Geo Photo", "Audience", "Mission Mapping", "PO-PSO Mapping"]
entry_vars = [tk.StringVar() for _ in labels]

# First Column
ttk.Label(main_frame, text="Event Number").grid(row=0, column=0, sticky='w', padx=5, pady=2)
event_number_entry = ttk.Entry(main_frame, textvariable=entry_vars[0])
event_number_entry.grid(row=1, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Event Name").grid(row=2, column=0, sticky='w', padx=5, pady=2)
event_name_entry = ttk.Entry(main_frame, textvariable=entry_vars[3])
event_name_entry.grid(row=3, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Event IC").grid(row=4, column=0, sticky='w', padx=5, pady=2)
event_ic_entry = ttk.Entry(main_frame, textvariable=entry_vars[5])
event_ic_entry.grid(row=5, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Date").grid(row=6, column=0, sticky='w', padx=5, pady=2)
date_entry = ttk.Entry(main_frame, textvariable=entry_vars[8])
date_entry.grid(row=7, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Event Type").grid(row=8, column=0, sticky='w', padx=5, pady=2)
event_type_entry = ttk.Entry(main_frame, textvariable=entry_vars[11])
event_type_entry.grid(row=9, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Report Doc").grid(row=10, column=0, sticky='w', padx=5, pady=2)
report_doc_entry = ttk.Entry(main_frame, textvariable=entry_vars[14])
report_doc_entry.grid(row=11, column=0, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Geo Photo").grid(row=12, column=0, sticky='w', padx=5, pady=2)
geo_photo_entry = ttk.Entry(main_frame, textvariable=entry_vars[17])
geo_photo_entry.grid(row=13, column=0, sticky='ew', padx=5, pady=2)


# Second Column
ttk.Label(main_frame, text="Attendees").grid(row=0, column=1, sticky='w', padx=5, pady=2)
attendees_entry = ttk.Entry(main_frame, textvariable=entry_vars[1])
attendees_entry.grid(row=1, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Resource Person").grid(row=2, column=1, sticky='w', padx=5, pady=2)
resource_person_entry = ttk.Entry(main_frame, textvariable=entry_vars[4])
resource_person_entry.grid(row=3, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Designation").grid(row=4, column=1, sticky='w', padx=5, pady=2)
designation_entry = ttk.Entry(main_frame, textvariable=entry_vars[6])
designation_entry.grid(row=5, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Address").grid(row=6, column=1, sticky='w', padx=5, pady=2)
address_entry = ttk.Entry(main_frame, textvariable=entry_vars[9])
address_entry.grid(row=7, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Funding").grid(row=8, column=1, sticky='w', padx=5, pady=2)
funding_entry = ttk.Entry(main_frame, textvariable=entry_vars[12])
funding_entry.grid(row=9, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Days").grid(row=10, column=1, sticky='w', padx=5, pady=2)
days_entry = ttk.Entry(main_frame, textvariable=entry_vars[15])
days_entry.grid(row=11, column=1, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Audience").grid(row=12, column=1, sticky='w', padx=5, pady=2)
audience_entry = ttk.Entry(main_frame, textvariable=entry_vars[18])
audience_entry.grid(row=13, column=1, sticky='ew', padx=5, pady=2)


# Third Column
ttk.Label(main_frame, text="Mission Mapping").grid(row=0, column=2, sticky='w', padx=5, pady=2)
mission_mapping_entry = ttk.Entry(main_frame, textvariable=entry_vars[2])
mission_mapping_entry.grid(row=1, column=2, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="PO-PSO Mapping").grid(row=2, column=2, sticky='w', padx=5, pady=2)
po_pso_mapping_entry = ttk.Entry(main_frame, textvariable=entry_vars[20])
po_pso_mapping_entry.grid(row=3, column=2, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Attendance Check").grid(row=4, column=2, sticky='w', padx=5, pady=2)
attendance_check_entry = ttk.Entry(main_frame, textvariable=entry_vars[7])
attendance_check_entry.grid(row=5, column=2, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Permission Docs").grid(row=6, column=2, sticky='w', padx=5, pady=2)
permission_docs_entry = ttk.Entry(main_frame, textvariable=entry_vars[10])
permission_docs_entry.grid(row=7, column=2, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="CO-PO Link").grid(row=8, column=2, sticky='w', padx=5, pady=2)
co_po_link_entry = ttk.Entry(main_frame, textvariable=entry_vars[13])
co_po_link_entry.grid(row=9, column=2, sticky='ew', padx=5, pady=2)

ttk.Label(main_frame, text="Remarks").grid(row=10, column=2, sticky='w', padx=5, pady=2)
remarks_entry = ttk.Entry(main_frame, textvariable=entry_vars[16])
remarks_entry.grid(row=11, column=2, sticky='ew', padx=5, pady=2)
# Text Areas
text_frame = ttk.Frame(main_frame)
text_frame.grid(row=14, column=0, columnspan=3, sticky='nsew', pady=15)

objective_text = tk.Text(text_frame, height=5, wrap='word')
outcome_text = tk.Text(text_frame, height=5, wrap='word')
course_contents_text = tk.Text(text_frame, height=5, wrap='word')

ttk.Label(text_frame, text="Objective of the Session", font=('Arial', 10, 'bold')).pack(anchor='w', pady=(0, 3))
objective_text.pack(fill='both', expand=True, padx=5, pady=(0, 10))
ttk.Label(text_frame, text="Outcome of the Session", font=('Arial', 10, 'bold')).pack(anchor='w', pady=(0, 3))
outcome_text.pack(fill='both', expand=True, padx=5, pady=(0, 10))
ttk.Label(text_frame, text="Course Contents", font=('Arial', 10, 'bold')).pack(anchor='w', pady=(0, 3))
course_contents_text.pack(fill='both', expand=True, padx=5, pady=(0, 10))

# Buttons
button_frame = ttk.Frame(input_tab)
button_frame.pack(fill='x', pady=15, padx=10)

ttk.Button(button_frame, text="Add Event", command=add_event).pack(side='left', padx=12)
ttk.Button(button_frame, text="Clear Fields", command=clear_fields).pack(side='left', padx=12)
ttk.Button(button_frame, text="Generate Report", command=generate_report).pack(side='left', padx=12)

# Documents & Images Tab - Improved UI
upload_frame = ttk.Frame(upload_tab)
upload_frame.pack(fill='both', expand=True, padx=15, pady=15)

# Create a container for the two main sections
container = ttk.Frame(upload_frame)
container.pack(fill='both', expand=True)

# Left side - Event Images
image_upload_frame = ttk.LabelFrame(container, text="Event Images", padding=10)
image_upload_frame.pack(side='left', fill='both', expand=True, padx=5, pady=5)

# Button to add images
ttk.Button(image_upload_frame, text="Add Event Images", command=add_images).pack(pady=5)

# Preview area for images
image_preview_frame = ttk.Frame(image_upload_frame)
image_preview_frame.pack(fill='both', expand=True, pady=10)

# Add a scrollbar for the preview
preview_canvas = tk.Canvas(image_upload_frame)
scrollbar = ttk.Scrollbar(image_upload_frame, orient="vertical", command=preview_canvas.yview)
scrollable_frame = ttk.Frame(preview_canvas)

scrollable_frame.bind(
    "<Configure>",
    lambda e: preview_canvas.configure(
        scrollregion=preview_canvas.bbox("all")
    )
)

preview_canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
preview_canvas.configure(yscrollcommand=scrollbar.set)

preview_canvas.pack(side="left", fill="both", expand=True)
scrollbar.pack(side="right", fill="y")

image_preview_frame = scrollable_frame  # Now our previews will go in the scrollable frame

# Right side - Documents
documents_frame = ttk.LabelFrame(container, text="Event Documents", padding=10)
documents_frame.pack(side='right', fill='both', padx=5, pady=5)

# Document upload sections
document_types = [
    ("Invitation Letter", "invitation"),
    ("Permission Document", "permission"),
    ("Certificate", "certificate"),
    ("CO-PO Mapping", "co_po"),
    ("Leaflet", "leaflet")
]

# Create variables and labels
invitation_path = tk.StringVar()
permission_path = tk.StringVar()
certificate_path = tk.StringVar()
co_po_path = tk.StringVar()
leaflet_path = tk.StringVar()

file_vars = {
    "invitation": invitation_path,
    "permission": permission_path,
    "certificate": certificate_path,
    "co_po": co_po_path,
    "leaflet": leaflet_path
}

file_labels = {}

for doc_type, doc_key in document_types:
    frame = ttk.Frame(documents_frame)
    frame.pack(fill='x', pady=5)

    ttk.Label(frame, text=f"{doc_type}:").pack(side='left', padx=5)
    btn = ttk.Button(frame, text="Upload",
                    command=lambda k=doc_key: upload_special_image(file_vars[k], k.replace("_", " ").title()))
    btn.pack(side='left', padx=5)

    lbl = ttk.Label(frame, text="No file selected", wraplength=200)
    lbl.pack(side='left', padx=5, fill='x', expand=True)
    file_labels[lbl] = file_vars[doc_key]

# Event List Tab
search_frame = ttk.Frame(list_tab)
search_frame.pack(fill='x', padx=10, pady=10)

search_var = tk.StringVar()
ttk.Label(search_frame, text="Search:").pack(side='left', padx=5)
search_entry = ttk.Entry(search_frame, textvariable=search_var)
search_entry.pack(side='left', fill='x', expand=True, padx=5)
ttk.Button(search_frame, text="Search", command=search_events).pack(side='left', padx=5)

tree = ttk.Treeview(list_tab, columns=("Event Number", "Event Name", "Date", "Resource Person", "Event Type", "Images"), show='headings')
for col in tree["columns"]:
    tree.heading(col, text=col)
tree.column("Event Number", width=100)
tree.column("Event Name", width=200)
tree.column("Date", width=100)
tree.column("Resource Person", width=150)
tree.column("Event Type", width=100)
tree.column("Images", width=80)
tree.pack(fill='both', expand=True, padx=10, pady=10)

button_frame = ttk.Frame(list_tab)
button_frame.pack(fill='x', pady=10)
ttk.Button(button_frame, text="Delete Selected Event", command=delete_event).pack(side='left', padx=10)
ttk.Button(button_frame, text="Generate Report", command=generate_report).pack(side='left', padx=10)

update_event_list()
update_file_labels() # Initialize file labels
root.mainloop()
